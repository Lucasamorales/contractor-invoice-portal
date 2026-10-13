"""Spreadsheet to database, with a report (R-SEED-01 .. R-SEED-13).

    uv run python -m app.scripts.seed --file "data/Data Invoices 1.xlsx"
    uv run python -m app.scripts.seed --file "..." --commit

Dry run is the default and prints what it would do. Nothing is written without
``--commit``, and the report is meant to be read by a person before that
happens: the spreadsheet is the least trustworthy input this system has.
"""

from __future__ import annotations

import argparse
import sys
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.db.models import (
    BankDetail,
    BankSource,
    BankState,
    Counter,
    Person,
    PersonStatus,
)
from app.db.models import PayoutRegion as DbPayoutRegion
from app.db.repo import get_person_by_email
from app.db.session import get_sessionmaker
from app.domain.seeding import COLUMNS, SeedRow, build_rows
from app.services.crypto import Cipher, field_aad

SHEET_NAME = "Data Invoices"

ENCRYPTED_FIELDS = ("account_number", "sort_code", "iban", "swift_bic")


@dataclass(slots=True)
class Outcome:
    """What the import did, or would do, for one row."""

    row: SeedRow
    action: str  # created | updated | skipped
    detail: str = ""


def read_rows(path: Path) -> list[dict[str, object]]:
    """Read the sheet into dicts keyed by the spreadsheet's own headings.

    ``data_only`` so a formula cell yields its value rather than its formula.
    """
    workbook = load_workbook(path, data_only=True, read_only=True)
    if SHEET_NAME not in workbook.sheetnames:
        raise SystemExit(
            f"{path} has no sheet named {SHEET_NAME!r}; found "
            f"{', '.join(workbook.sheetnames)}"
        )
    sheet = workbook[SHEET_NAME]
    rows = sheet.iter_rows(values_only=True)

    try:
        header = [str(cell or "").strip() for cell in next(rows)]
    except StopIteration as exc:
        raise SystemExit(f"{path} is empty") from exc

    missing = [column for column in COLUMNS if column not in header]
    if missing:
        raise SystemExit(f"missing expected columns: {', '.join(missing)}")

    out: list[dict[str, object]] = []
    for values in rows:
        record = dict(zip(header, values, strict=False))
        # The sheet is a table with a thousand blank rows below the data.
        if not str(record.get("EMAIL") or "").strip():
            continue
        out.append(record)
    workbook.close()
    return out


def report(rows: list[SeedRow], outcomes: list[Outcome], *, committed: bool) -> int:
    """Print the report and return the number of rows needing attention."""
    print()
    print(f"{'COMMITTED' if committed else 'DRY RUN'} — {len(rows)} rows read")
    print("=" * 72)

    for outcome in outcomes:
        row = outcome.row
        print()
        print(f"row {row.source_row}: {row.display_name or '(no name)'} <{row.email}>")
        note = f" — {outcome.detail}" if outcome.detail else ""
        print(f"  action        {outcome.action}{note}")
        print(f"  team          {row.team or '—'}")
        print(
            f"  payout        {row.payout_region.value if row.payout_region else '—'}"
        )
        print(f"  legal name    {row.legal_name or '—'}")
        print(f"  next invoice  {row.next_number if row.next_number else '—'}")

        bank = row.bank.masked()
        print(
            f"  bank          name={bank['account_name'] or '—'} "
            f"bank={bank['bank_name'] or '—'}"
        )
        print(
            f"                sort={bank['sort_code'] or '—'} "
            f"acct={bank['account_number'] or '—'} "
            f"iban={bank['iban'] or '—'} bic={bank['swift_bic'] or '—'}"
        )

        for error in row.errors:
            print(f"  ERROR         {error}")
        for warning in row.warnings:
            print(f"  warning       {warning}")

    needing_attention = sum(1 for row in rows if row.needs_attention)
    blocked = sum(1 for row in rows if not row.importable)

    print()
    print("=" * 72)
    print(
        f"{len(rows)} rows, {blocked} blocked by errors, "
        f"{needing_attention} needing human attention"
    )
    if not committed:
        print("Nothing was written. Re-run with --commit to apply.")
    return needing_attention


def apply_row(
    session: Session,
    row: SeedRow,
    cipher: Cipher | None,
    *,
    kudos_ineligible: tuple[str, ...] = (),
) -> Outcome:
    """Create or update one person, their counter and their bank row.

    Idempotent on email (R-SEED-12): a second run updates rather than
    duplicating, and never replaces a bank row the person entered themselves
    with one from the spreadsheet.
    """
    person = get_person_by_email(session, row.email)
    action = "updated"
    if person is None:
        person = Person(email=row.email, display_name=row.display_name)
        session.add(person)
        action = "created"

    person.display_name = row.display_name
    person.legal_name = row.legal_name
    person.team = row.team
    # R-KUD-20. The same list sign-in applies (Settings.kudos_ineligible_emails).
    person.kudos_eligible = row.email.strip().lower() not in kudos_ineligible
    # R-SEED-11: imported people are active with an unconfirmed profile. They
    # still cannot invoice until Robin sets a fee (R-WIZ-09).
    person.status = PersonStatus.ACTIVE
    person.profile_confirmed_at = None
    if row.payout_region is not None:
        person.payout_region = DbPayoutRegion(row.payout_region.value)
    session.flush()

    counter = session.get(Counter, person.id)
    if counter is None:
        session.add(Counter(person_id=person.id, next_number=row.next_number or 1))
    else:
        # R-SEED-09 seeds the sequence; it must never wind a live counter back.
        counter.next_number = max(counter.next_number, row.next_number or 1)

    detail = _apply_bank(session, person, row, cipher)
    return Outcome(row=row, action=action, detail=detail)


def _apply_bank(
    session: Session, person: Person, row: SeedRow, cipher: Cipher | None
) -> str:
    """Write the imported bank row, unless the person has entered their own."""
    existing = [
        detail
        for detail in person.bank_details
        if detail.source in (BankSource.SELF, BankSource.ADMIN_PROPOSED)
    ]
    if existing:
        return "bank row left alone: the person has entered their own"

    already_imported = [
        detail for detail in person.bank_details if detail.source == BankSource.IMPORT
    ]
    if already_imported:
        return "bank row already imported"

    if cipher is None:
        return "bank row not written (no encryption key configured)"

    values = row.bank.as_dict()
    detail = BankDetail(
        person_id=person.id,
        account_name=values["account_name"],
        bank_name=values["bank_name"],
        source=BankSource.IMPORT,
        # R-SEED-10: unverified regardless of how clean it looked. Nothing in
        # the spreadsheet is treated as true, so an imported row cannot produce
        # an invoice until its owner confirms it (R-BANK-02).
        state=BankState.UNVERIFIED,
        verified_at=None,
        encryption_key_id=cipher.key_id,
    )
    for name in ENCRYPTED_FIELDS:
        value = values[name]
        if value:
            setattr(
                detail,
                f"{name}_enc",
                cipher.encrypt(value, aad=field_aad(str(person.id), name)),
            )
    session.add(detail)
    return "bank row imported, unverified"


def main(argv: list[str] | None = None) -> int:
    """Command-line entry point: read the spreadsheet, print the report, write only
    with ``--commit``.
    """
    parser = argparse.ArgumentParser(
        prog="seed", description="Import the invoice spreadsheet."
    )
    parser.add_argument("--file", required=True, type=Path, help="the .xlsx to read")
    parser.add_argument(
        "--commit",
        action="store_true",
        help="write to the database; without this nothing is written",
    )
    args = parser.parse_args(argv)

    if not args.file.exists():
        raise SystemExit(f"no such file: {args.file}")

    rows = build_rows(read_rows(args.file))

    if not args.commit:
        outcomes = [
            Outcome(row=row, action="would import" if row.importable else "blocked")
            for row in rows
        ]
        attention = report(rows, outcomes, committed=False)
        return 1 if attention else 0

    blocked = [row for row in rows if not row.importable]
    if blocked:
        report(
            rows, [Outcome(row=row, action="blocked") for row in rows], committed=False
        )
        print()
        print(f"Refusing to commit: {len(blocked)} rows have errors.")
        return 2

    from app.config import get_settings

    settings = get_settings()
    cipher = (
        Cipher.from_settings(settings)
        if settings.bank_encryption_key.get_secret_value()
        else None
    )
    if cipher is None:
        print("Refusing to commit: no bank encryption key configured (R-SEC-03).")
        return 2

    factory = get_sessionmaker()
    with factory() as session, session.begin():
        outcomes = [
            apply_row(
                session,
                row,
                cipher,
                kudos_ineligible=settings.kudos_ineligible_emails,
            )
            for row in rows
        ]

    report(rows, outcomes, committed=True)
    return 0


if __name__ == "__main__":
    sys.exit(main())
